# openpyxl Formatting Examples

## Purpose
This file defines the standard formatting styles and patterns the copilot agent must apply when generating Excel output. Always use these definitions — do not invent alternate styles.

---

## Required Imports
Always include all of these at the top of any script that writes Excel files:

```python
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
```

---

## Standard Style Definitions

Define these at the top of the script, before writing any data:

```python
# Fonts
header_font = Font(bold=True, size=11)
body_font = Font(size=11)

# Fill
header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

# Borders
thin_side = Side(border_style="thin", color="000000")
thin_border = Border(
    top=thin_side,
    left=thin_side,
    right=thin_side,
    bottom=thin_side
)

# Alignment
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")

# Number formats
currency_format = "$#,##0.00"
percent_format = "0.0%"
date_format = "yyyy-mm-dd"
integer_format = "#,##0"
```

---

## Writing a DataFrame to a Sheet

Always write DataFrames using `dataframe_to_rows` with `index=False, header=True`:

```python
# ✅ CORRECT
for row in dataframe_to_rows(df, index=False, header=True):
    ws.append(row)

# ❌ WRONG — writing raw values loses structure and requires manual column tracking
ws.append(list(df.columns))
for _, row in df.iterrows():
    ws.append(list(row))
```

---

## Formatting the Header Row

Always apply header formatting after writing data. The header is always row 1:

```python
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = center_align
```

---

## Applying Number Formats to Data Rows

After writing data, apply number formats by column. Use `get_column_letter` to reference columns by index:

```python
# Example: format column 3 (revenue) as currency, column 4 (rate) as percent
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
    for cell in row:
        cell.number_format = currency_format

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=4):
    for cell in row:
        cell.number_format = percent_format
```

---

## Auto-Fitting Column Widths

openpyxl does not auto-fit columns natively. Use this helper:

```python
def autofit_columns(ws, min_width=10, max_width=40):
    for col in ws.columns:
        max_len = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col
        )
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)

autofit_columns(ws)
```

---

## Inserting a Chart Image

Save the chart as a PNG first (see Matplotlib style guide), then embed it in the sheet:

```python
img = Image("chart.png")
img.anchor = "H2"          # top-left cell for the image
ws.add_image(img)
```

---

## Complete Workbook Example

```python
from datetime import date
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

# Style definitions
header_font = Font(bold=True, size=11)
header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
thin_side = Side(border_style="thin", color="000000")
thin_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
center_align = Alignment(horizontal="center", vertical="center")
currency_format = "$#,##0.00"

def autofit_columns(ws, min_width=10, max_width=40):
    for col in ws.columns:
        max_len = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col
        )
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)

# Build workbook
wb = Workbook()
ws = wb.active
ws.title = "Summary"

# Write data
for row in dataframe_to_rows(report, index=False, header=True):
    ws.append(row)

# Format header
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.border = thin_border
    cell.alignment = center_align

# Format revenue column (column 3) as currency
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=3, max_col=3):
    for cell in row:
        cell.number_format = currency_format

# Auto-fit columns
autofit_columns(ws)

# Embed chart
img = Image("chart.png")
img.anchor = "H2"
ws.add_image(img)

# Save with dynamic filename
today = date.today()
filename = f"revenue_report_{today.strftime('%Y_%m')}.xlsx"
wb.save(filename)
print(f"Saved: {filename}")
```
